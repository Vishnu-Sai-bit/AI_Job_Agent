import re
from typing import Tuple, Optional
import pymupdf  # PyMuPDF modern import
import docx

def extract_text_from_pdf(file_path: str) -> Tuple[str, Optional[str]]:
    """
    Extracts text from a PDF file using PyMuPDF.
    Returns (extracted_text, error_message).
    """
    text_chunks = []
    try:
        doc = pymupdf.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text:
                text_chunks.append(page_text)
        doc.close()
        
        full_text = "\n".join(text_chunks).strip()
        if not full_text:
            return "", "PDF contains no extractable text. It may be a scanned image or empty document."
        
        return clean_extracted_text(full_text), None
    except Exception as e:
        return "", f"Error reading PDF: {str(e)}"

def extract_text_from_docx(file_path: str) -> Tuple[str, Optional[str]]:
    """
    Extracts text from a DOCX file using python-docx.
    Returns (extracted_text, error_message).
    """
    text_chunks = []
    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            if p.text.strip():
                text_chunks.append(p.text.strip())
        
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_chunks.append(" | ".join(row_cells))
                    
        full_text = "\n".join(text_chunks).strip()
        if not full_text:
            return "", "DOCX contains no extractable text."
            
        return clean_extracted_text(full_text), None
    except Exception as e:
        return "", f"Error reading DOCX: {str(e)}"

def extract_text(file_path: str, file_type: str) -> Tuple[str, Optional[str]]:
    """
    Generic dispatcher for extracting text by file type.
    """
    ext = file_type.lower().strip(".")
    if ext == "pdf":
        return extract_text_from_pdf(file_path)
    elif ext == "docx":
        return extract_text_from_docx(file_path)
    else:
        return "", f"Unsupported file type: {ext}"

def clean_extracted_text(text: str) -> str:
    """
    Normalizes spacing, removes redundant empty lines, and standardizes bullet points.
    """
    text = text.replace("\xa0", " ")
    text = re.sub(r"[\u2022\u2023\u25E6\u2043\u2219]", "*", text)
    text = re.sub(r"[\u2013\u2014]", "-", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
